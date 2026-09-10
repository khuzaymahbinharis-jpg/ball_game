from __future__ import annotations

from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "reports"
ASSET_DIR = REPORT_DIR / "assets"
OUTPUT = REPORT_DIR / "track_the_game_final_report.docx"

NAVY = "17365D"
BLUE = "2F75B5"
CYAN = "00A6D6"
PALE = "EAF2F8"
LIGHT = "F4F6F7"
MID = "D6E4F0"
TEXT = RGBColor(31, 41, 55)


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top=60, start=75, bottom=60, end=75) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + edge
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color="B8C7D9", size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_col_width(cell, width_inches: float) -> None:
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_text(paragraph, text: str, *, bold=False, size=8.7, color=TEXT, italic=False) -> None:
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.color.rgb = color


def add_body(doc: Document, text: str, *, first_line=True, space_after=3.5) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.04
    if first_line:
        p.paragraph_format.first_line_indent = Inches(0.18)
    add_text(p, text)


def add_heading(doc: Document, text: str, level=1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(5 if level == 1 else 3)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Aptos Display"
    run.font.color.rgb = RGBColor.from_string(NAVY if level == 1 else BLUE)
    run.bold = True
    run.font.size = Pt(13 if level == 1 else 10.5)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = False
    add_text(p, text, size=7.6, color=RGBColor(75, 85, 99), italic=True)


def add_page_number(section) -> None:
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    add_text(p, "Zeta Solutions  |  Track the Game  |  ", size=7.5, color=RGBColor(95, 99, 104))
    run = p.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def prepare_images() -> tuple[Path, Path]:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    success_src = ROOT / "outputs" / "previews" / "gemini_3_8_wide_new.png"
    success_dst = ASSET_DIR / "final_tracking_frame.jpg"
    with Image.open(success_src) as im:
        im.convert("RGB").resize((1280, 720), Image.Resampling.LANCZOS).save(success_dst, quality=90)

    comparison_src = ROOT / "outputs" / "tracker_comparison_contact_2.jpg"
    comparison_dst = ASSET_DIR / "player_cv_ablation.jpg"
    with Image.open(comparison_src) as im:
        # Bottom pair is a wide-play frame using the same saved VLM detections.
        crop = im.crop((0, 1325, im.width, im.height))
        crop.save(comparison_dst, quality=92)
    return success_dst, comparison_dst


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.48)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.22)
    add_page_number(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(8.7)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(3.5)

    for name in ("Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style.font.color.rgb = RGBColor.from_string(NAVY)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("Track the Game")
    r.font.name = "Aptos Display"
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(NAVY)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    add_text(p, "A hybrid VLM and classical vision pipeline for sports tracking", size=10.2, color=RGBColor.from_string(BLUE))

    bar = doc.add_table(rows=1, cols=2)
    bar.autofit = False
    bar.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_borders(bar, color=NAVY, size="0")
    set_col_width(bar.cell(0, 0), 5.55)
    set_col_width(bar.cell(0, 1), 1.85)
    set_cell_shading(bar.cell(0, 0), NAVY)
    set_cell_shading(bar.cell(0, 1), CYAN)
    for cell in bar.rows[0].cells:
        set_cell_margins(cell, top=15, bottom=15)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    add_text(p, "Zeta Solutions internship technical report  •  September 2026", size=7.8, color=RGBColor(90, 99, 109))


def add_pipeline(doc: Document) -> None:
    labels = ["Video", "Anchor\nframes", "Ruler", "VLM", "Structured\ndetections", "Association\nand tracking", "Annotated\nvideo"]
    table = doc.add_table(rows=1, cols=len(labels))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="FFFFFF", size="3")
    widths = [0.72, 0.92, 0.68, 0.62, 1.08, 1.18, 0.98]
    for idx, (cell, label, width) in enumerate(zip(table.rows[0].cells, labels, widths)):
        set_col_width(cell, width)
        set_cell_shading(cell, NAVY if idx in (0, 3, 6) else (BLUE if idx in (1, 4) else MID))
        set_cell_margins(cell, top=75, bottom=75, start=35, end=35)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        add_text(p, label, bold=True, size=7.3, color=RGBColor(255, 255, 255) if idx in (0, 1, 3, 4, 6) else TEXT)
    add_caption(doc, "Figure 1  Final data flow  Semantic reasoning occurs only in the VLM stage")


def add_page_one(doc: Document, success_image: Path) -> None:
    add_title(doc)
    add_heading(doc, "Method and final pipeline")
    add_body(doc, "The task is to add stable player, team, ball, and possession overlays to 30-second sports clips without running an expensive semantic model on all 900 frames. The final design samples ruler-augmented anchor frames, asks the VLM for schema-validated normalized coordinates and labels, then uses those anchors to constrain a faster local tracker. This division preserves semantic capability while keeping the number of API calls bounded.")
    add_pipeline(doc)
    add_body(doc, "The ruler supplies a consistent spatial reference when players are small. For players, Hungarian assignment minimizes a joint cost over foot-point distance, motion, overlap, size change, and a soft team-history penalty. Soft history is preferable to a hard team constraint because a single VLM team error should not force a new identity. Tracks persist through short missing-anchor gaps and move among confirmed, uncertain, and lost states. Sparse Lucas–Kanade features update each VLM-initialized player between anchors; partial-affine camera estimation removes broadcast pan and modest zoom from apparent player motion. Scene-cut detection resets visual state, and close-up suppression avoids treating a single broadcast portrait as a playable wide shot.")
    add_body(doc, "The ball uses a separate VLM-initialized track because its scale and acceleration differ from players. Hungarian gating compares each new ball anchor with the motion prediction, optical flow coasts through short gaps, and a Kalman filter smooths the accepted trajectory. Possession acts only as a weak prior. Pillow renders filled team-colour ellipses at feet, a ball marker, and possession indicator onto the final muted video.")

    boundary = doc.add_table(rows=1, cols=2)
    boundary.alignment = WD_TABLE_ALIGNMENT.CENTER
    boundary.autofit = False
    set_table_borders(boundary, color="B8C7D9", size="4")
    for cell, title, body, color in (
        (boundary.cell(0, 0), "Semantic boundary", "VLM: player and team labels, ball, possession, and normalized locations.", "E8F1FA"),
        (boundary.cell(0, 1), "Nonsemantic boundary", "Classical CV: association, VLM-initialized optical flow, camera motion, cuts, prediction, and rendering.", "E9F7FA"),
    ):
        set_col_width(cell, 3.65)
        set_cell_shading(cell, color)
        set_cell_margins(cell, top=65, bottom=65)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(1)
        add_text(p, title + "\n", bold=True, size=8.2, color=RGBColor.from_string(NAVY))
        add_text(p, body, size=7.8)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    p.add_run().add_picture(str(success_image), width=Inches(6.55))
    add_caption(doc, "Figure 2  Representative Gemini 3.8 Flash frame with filled team markers and ball or possession emphasis")


def add_experiment_table(doc: Document) -> None:
    rows = [
        ("2 FPS baseline", "Greedy and linear", "$0.161052", "75.30 s", "34 created, 24 expired, 41 lost"),
        ("2 FPS persistent", "Hungarian and soft team", "same anchors", "88.52 s", "27 created, 17 expired, 36 lost"),
        ("Player CV ablation", "LK, camera, cuts, confidence", "$0 new", "106.54 s", "Better visual following; more local work"),
        ("1 FPS", "31 anchors", "$0.083161", "150.24 s", "27 valid; quality rejected"),
        ("3 FPS multisport", "Gemini 3.8, five clips", "$0.618–0.882", "979.72 s*", "All clips stayed below $1"),
    ]
    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = ["Experiment", "Main change", "API cost", "Runtime", "Main finding"]
    widths = [1.15, 1.48, 0.85, 0.83, 2.8]
    for i, (cell, text, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
        set_col_width(cell, width)
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        add_text(p, text, bold=True, size=7.4, color=RGBColor(255, 255, 255))
    set_repeat_table_header(table.rows[0])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, (cell, text, width) in enumerate(zip(cells, row, widths)):
            set_col_width(cell, width)
            set_cell_shading(cell, "FFFFFF" if ridx % 2 == 0 else LIGHT)
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_text(p, text, size=7.2)
    set_table_borders(table)
    add_caption(doc, "Table 1  Measured development experiments  *Basketball 3 FPS total is API batch plus combined local processing")


def add_page_two(doc: Document, comparison_image: Path) -> None:
    doc.add_page_break()
    add_heading(doc, "Experiments successes and failures")
    add_body(doc, "The initial system used sparse VLM anchors, greedy nearest-neighbour assignment, a hard team constraint, linear player interpolation, and simple ball interpolation. It produced a complete video, but visual review exposed identity fragmentation, markers gliding along straight paths while players changed direction, camera pans being mistaken for object motion, frequent ball loss, and gaps after schema-invalid responses.")
    add_experiment_table(doc)
    add_heading(doc, "Controlled association comparison", level=2)
    add_body(doc, "Test 3 reused the same 61 Gemini 3.1 Flash Lite anchor requests for both trackers, of which 49 were schema valid. Replacing greedy matching with Hungarian assignment, soft team history, persistence, and the dedicated ball path reduced created operational IDs from 34 to 27, expired tracks from 24 to 17, and lost events from 41 to 36. These are machine lifecycle diagnostics, not manually verified ID-switch counts; the repository contains no identity ground truth.")
    add_heading(doc, "Between anchor player CV", level=2)
    add_body(doc, "Test 4 again reused the saved detections, so it added no API cost. Sparse LK and camera compensation followed curved or accelerating movement more naturally than straight-line interpolation in representative frame review. The machine log recorded 859 successful and 40 failed camera estimates, three cut resets at frames 15, 369, and 442, and 392 uncertain-track frames. The tradeoff was local processing: 31.10 s for the current improved tracker versus 51.03 s for the CV variant, including an 18.46 s player-CV pass and 18.22 s rendering.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    p.add_run().add_picture(str(comparison_image), width=Inches(6.55))
    add_caption(doc, "Figure 3  Same-frame ablation  Linear state propagation at left and sparse-LK correction at right")
    add_heading(doc, "Sampling frequency", level=2)
    add_body(doc, "At 1 FPS, 31 calls produced 27 valid responses and four schema failures for $0.08316075, a 48.36% API-cost reduction from the controlled 2 FPS run. This was a useful negative result: longer correction gaps made uncertainty and ball loss more visible, while local processing reached 112.99 s and total time 150.24 s. We therefore rejected 1 FPS as the primary setting despite its price.")


def add_model_table(doc: Document) -> None:
    rows = [
        ("Gemini 3.8 Flash", "98.36%", "8.12", "44 / 41", "$0.371564", "7.57 s"),
        ("Gemini 3.7 Flash", "98.36%", "8.07", "48 / 44", "$0.383148", "10.78 s"),
        ("Gemini 3.1 Flash Lite", "80.33%", "7.41", "40 / 37", "$0.161052", "5.65 s"),
        ("Qwen3 VL 235B", "100.00%", "0.00", "0 / 0", "$0.055579", "4.26 s"),
        ("Qwen3 VL 30B", "83.61%", "7.78", "50 / 50", "$0.053971+", "11.37 s"),
        ("Seed 2.1 Turbo", "68.85%", "7.71", "21 / 22", "$0.268502", "21.36 s"),
        ("GLM 5.3 Flash", "18.03%", "7.27", "7 / 6", "$0.006618+", "26.63 s"),
    ]
    table = doc.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = ["Model", "Schema valid", "Players per valid", "Ball / poss.", "Cost", "Mean latency"]
    widths = [1.55, 0.85, 1.02, 0.85, 0.9, 1.0]
    for cell, text, width in zip(table.rows[0].cells, headers, widths):
        set_col_width(cell, width)
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell, top=55, bottom=55)
        p = cell.paragraphs[0]
        add_text(p, text, bold=True, size=7.1, color=RGBColor(255, 255, 255))
    set_repeat_table_header(table.rows[0])
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for cell, text, width in zip(cells, row, widths):
            set_col_width(cell, width)
            set_cell_shading(cell, "EAF2F8" if ridx == 0 else ("FFFFFF" if ridx % 2 == 0 else LIGHT))
            set_cell_margins(cell, top=45, bottom=45)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_text(p, text, bold=(ridx == 0), size=7.0)
    set_table_borders(table)
    add_caption(doc, "Table 2  Two FPS model bake-off over 61 planned anchors  + denotes incomplete cost accounting")


def add_callout(doc: Document, title: str, body: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color=CYAN, size="7")
    cell = table.cell(0, 0)
    set_cell_shading(cell, "E9F8FC")
    set_cell_margins(cell, top=70, bottom=70, start=100, end=100)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_text(p, title + "  ", bold=True, size=8.2, color=RGBColor.from_string(NAVY))
    add_text(p, body, size=8.0)


def add_page_three(doc: Document) -> None:
    doc.add_page_break()
    add_heading(doc, "Analysis quality cost and speed")
    add_heading(doc, "Model selection", level=2)
    add_model_table(doc)
    add_body(doc, "Gemini 3.8 Flash was selected on the measured tradeoff rather than model size. In the bake-off it matched Gemini 3.7 Flash at 98.36% schema validity, cost 3.0% less, and returned in 7.57 s on average instead of 10.78 s. It also averaged slightly more players per valid anchor, although 3.7 returned ball and possession fields more often. Qwen3 VL 235B was cheap and schema-valid but returned no players, ball, or possession under this schema, so schema validity alone was misleading. Qwen3 VL 30B returned many semantic fields but had only 83.61% valid anchors and incomplete cost accounting. Foot-point error, team accuracy, ball accuracy, and possession accuracy were not manually annotated for any model, so the table reports completeness proxies rather than detection accuracy.")
    add_heading(doc, "Why the hybrid tracker helped", level=2)
    add_body(doc, "Hungarian assignment considers all players jointly, avoiding locally attractive greedy matches that block better global pairings. Soft team history tolerates occasional colour or jersey mistakes without discarding identity. Camera compensation separates global broadcast motion from local player motion, and sparse LK follows nonlinear displacement between semantic corrections. These methods still fail under occlusion, overlapping bodies, small or featureless players, strong rotation or zoom, and long gaps caused by invalid VLM output. The dedicated ball filter is necessary because a small ball can move farther between anchors than a player and may be blurred or hidden; gating and Kalman smoothing reduce implausible jumps but cannot recover an unobserved ball reliably.")
    add_heading(doc, "Final cost and latency", level=2)
    add_body(doc, "The final cross-sport evaluation used Gemini 3.8 Flash at 3 FPS: 91 calls for each 30-second clip. Across football, basketball, and volleyball, 420 of 455 responses were valid and per-video API cost ranged from $0.617863 to $0.882180. All five clips therefore met the <$1 API-cost requirement. The basketball run cost $0.617863, used 82 valid anchors, spent 404.54 s in the API batch, and 575.18 s in combined local tracking and rendering, for a derived 979.72 s total. The latest runner did not log rendering separately; the controlled CV ablation measured 18.22 s rendering within 51.03 s local time.")
    add_callout(doc, "Latency finding", "The system does not meet the ideal <15 s or accepted ≈25 s target. API response time and CPU optical-flow processing are both bottlenecks; recent 3–4 FPS runs were minutes, not seconds.")
    add_heading(doc, "Conclusion", level=2)
    add_body(doc, "The experiments support a hybrid design: sparse VLM semantics corrected by persistent, camera-aware local tracking. More anchors improve recovery but raise cost; too few anchors make every schema failure more damaging. Gemini 3.8 Flash provides the best tested reliability, cost, and latency balance, and 3 FPS generalized below $1 per 30-second clip across five sports videos. The remaining priorities are manual ground-truth evaluation, faster batched inference, and optimized local tracking; without those, visual quality is promising but neither identity accuracy nor production latency is established.")


def set_core_properties(doc: Document) -> None:
    props = doc.core_properties
    props.title = "Track the Game"
    props.subject = "Final technical report for the Zeta Solutions internship project"
    props.author = "Zeta Solutions internship project"
    props.keywords = "sports tracking, VLM, Hungarian assignment, optical flow, OpenRouter"


def main() -> None:
    success_image, comparison_image = prepare_images()
    doc = Document()
    style_document(doc)
    set_core_properties(doc)
    add_page_one(doc, success_image)
    add_page_two(doc, comparison_image)
    add_page_three(doc)
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
